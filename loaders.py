"""Turn PDF, DOCX, PPTX and image files into LangChain Documents."""
from pathlib import Path
from langchain_core.documents import Document

SUPPORTED = ["pdf", "docx", "pptx", "png", "jpg", "jpeg"]
_ocr = None


def _ocr_text(path: str) -> str:
    global _ocr
    if _ocr is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr = RapidOCR()
    result, _ = _ocr(path)
    return "\n".join(line[1] for line in (result or []))


def load_file(path: str, name: str | None = None) -> list[Document]:
    p = Path(path)
    name = name or p.name
    ext = p.suffix.lower().lstrip(".")
    pages: list[tuple[int | None, str]] = []

    if ext == "pdf":
        from pypdf import PdfReader
        pages = [(i, pg.extract_text() or "") for i, pg in enumerate(PdfReader(path).pages, 1)]
    elif ext == "docx":
        import docx
        d = docx.Document(path)
        text = "\n".join(par.text for par in d.paragraphs)
        for table in d.tables:
            for row in table.rows:
                text += "\n" + " | ".join(c.text for c in row.cells)
        pages = [(None, text)]
    elif ext == "pptx":
        from pptx import Presentation
        for i, slide in enumerate(Presentation(path).slides, 1):
            text = "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
            pages.append((i, text))
    elif ext in ("png", "jpg", "jpeg"):
        pages = [(None, _ocr_text(path))]
    else:
        raise ValueError(f"Unsupported file type: .{ext}")

    return [Document(page_content=t, metadata={"source": name, "page": n})
            for n, t in pages if t.strip()]
